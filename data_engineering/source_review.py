import json,hashlib,time
from backend.database.postgres import connection
from pathlib import Path
from openpyxl import load_workbook
import argparse
parser=argparse.ArgumentParser()
parser.add_argument('--source-dir',required=True)
parser.add_argument('--staging',default='staging',choices=['staging'])
parser.add_argument('--report',required=True)
args=parser.parse_args()
base=Path(args.source_dir)
def val(v):
    if v is None:return None
    if hasattr(v,'isoformat'):return v.isoformat()[:10]
    return str(v).strip()
with connection('staging') as db:
    if db.execute('SELECT 1 FROM inspection_source LIMIT 1').fetchone():
        raise ValueError('Staging data already exists. Preserve it and use a separate database for a new import.')
    manifest=[]
    for name,kind in [('e4832250-43dc-4c0c-982b-b31f87205869.xlsx','inspections'),('e00adefd-b9ce-414f-ae04-21c950fb5f27.xlsx','citations')]:
        wb=load_workbook(base/name,read_only=True,data_only=True);ws=wb.active;rows=ws.iter_rows(values_only=True);header=next(rows);batch=[];count=0
        for number,row in enumerate(rows,2):
            if not any(v is not None for v in row):continue
            r=dict(zip(header,row));r={k:val(v) for k,v in r.items()}
            if kind=='inspections':
                batch.append((number,r['Inspection ID'],r['FEI Number'],r['Legal Name'],r['City'],r['State'],r['Country/Area'],r['Fiscal Year'],r['Inspection End Date'],r['Classification'],r['Posted Citations'],r['Project Area'],r['Product Type'],json.dumps(r)))
            else:batch.append((number,r['Inspection ID'],r['FEI Number'],r['Legal Name'],r['Inspection End Date'],r['Program Area'],r['Act/CFR Number'],r['Short Description'],r['Long Description']))
            count+=1
            if len(batch)>=5000:
                db.cursor().executemany('INSERT INTO '+('inspection_source VALUES ('+','.join(['%s']*14)+')' if kind=='inspections' else 'citation_source VALUES ('+','.join(['%s']*9)+')'),batch);batch=[]
        if batch:db.cursor().executemany('INSERT INTO '+('inspection_source VALUES ('+','.join(['%s']*14)+')' if kind=='inspections' else 'citation_source VALUES ('+','.join(['%s']*9)+')'),batch)
        wb.close()
        manifest.append({'file':name,'kind':kind,'rows':count,'sha256':hashlib.sha256((base/name).read_bytes()).hexdigest()})
        print(kind,count,flush=True)
    summaries=[]
    for year,name in [(2023,'Inspection_Observations_FY23_0.xlsx'),(2024,'inspection_observations_fy24.xlsx'),(2025,'inspection_observations_fiscal_year_2025_0.xlsx')]:
        wb=load_workbook(base/name,read_only=True,data_only=True);count=0;sheet_counts={}
        for ws in wb:
            if ws.title=='Summary':
                summaries.append({'year':year,'rows':[[val(v) for v in row] for row in ws.iter_rows(values_only=True)]});continue
            rows=ws.iter_rows(values_only=True);header=next(rows);n=0
            for source_row,row in enumerate(rows,2):
                if len(row)<6 or row[1] is None:continue
                if not isinstance(row[5],(int,float)):continue
                db.execute('INSERT INTO annual_source VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)',(year,ws.title,source_row,*[val(v) for v in row[:5]],int(row[5])))
                n+=1;count+=1
            sheet_counts[ws.title]=n
        wb.close()
        manifest.append({'file':name,'kind':'annual_frequency','year':year,'rows':count,'sheet_counts':sheet_counts,'sha256':hashlib.sha256((base/name).read_bytes()).hexdigest()})
    def query(sql):return [dict(row) for c in [db.execute(sql)] for row in c.fetchall()]
    report={'sources':manifest,'annual_summaries':summaries,
    'inspection_stats':query('SELECT COUNT(*) rows,COUNT(DISTINCT id) inspections,COUNT(DISTINCT fei) sites,MIN(date) min_date,MAX(date) max_date FROM inspection_source'),
    'classification':query('SELECT classification,COUNT(*) n FROM inspection_source GROUP BY classification'),
    'posted_flags':query('SELECT posted,COUNT(*) n FROM inspection_source GROUP BY posted'),
    'conflicts':query('SELECT COUNT(*) ids FROM (SELECT id FROM inspection_source GROUP BY id HAVING COUNT(DISTINCT fei)>1 OR COUNT(DISTINCT date)>1)'),
    'multi_classification':query('SELECT COUNT(*) ids FROM (SELECT id FROM inspection_source GROUP BY id HAVING COUNT(DISTINCT classification)>1)'),
    'duplicate_examples':query("SELECT id,COUNT(*) n,string_agg(DISTINCT classification,',') classifications,string_agg(DISTINCT project,',') projects FROM inspection_source GROUP BY id HAVING COUNT(*)>1 LIMIT 5"),
    'citation_stats':query('SELECT COUNT(*) rows,COUNT(DISTINCT id) inspections FROM citation_source'),
    'citation_join':query('SELECT COUNT(*) matched_rows,COUNT(DISTINCT c.id) matched_inspections FROM citation_source c WHERE EXISTS(SELECT 1 FROM inspection_source i WHERE i.id=c.id AND i.fei=c.fei AND i.date=c.date)'),
    'orphan_examples':query('SELECT id,fei,date FROM citation_source c WHERE NOT EXISTS(SELECT 1 FROM inspection_source i WHERE i.id=c.id AND i.fei=c.fei AND i.date=c.date) LIMIT 10'),
    'annual_frequency':query('SELECT year,COUNT(*) citation_categories,SUM(frequency) frequency_sum FROM annual_source GROUP BY year'),
    'duplicate_citations':query('SELECT SUM(n-1) extras FROM (SELECT COUNT(*) n FROM citation_source GROUP BY id,fei,date,program,reference,short,long HAVING COUNT(*)>1)')}
    Path(args.report).write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='annual_summaries'},indent=2))
