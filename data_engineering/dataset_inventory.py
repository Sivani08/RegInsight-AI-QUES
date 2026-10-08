"""Inventory all source workbooks and prepared artifacts without treating outputs as new inputs."""
import hashlib,json
from collections import Counter
from pathlib import Path
import pyarrow.parquet as pq
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
def source_paths(root=ROOT):
    """Resolve operator-supplied paths against the project, preserving absolute paths."""
    root=Path(root);manifest=root/'config/dataset_sources.json'
    if not manifest.exists():return []
    paths=[Path(p) for p in json.loads(manifest.read_text(encoding='utf-8-sig'))['workbooks']]
    return [p if p.is_absolute() else root/p for p in paths]
def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''):h.update(chunk)
    return h.hexdigest()
def inventory(root=ROOT):
    root=Path(root);out=root/'data/intelligence';out.mkdir(parents=True,exist_ok=True)
    cached_path=out/'inventory.json';cached=json.loads(cached_path.read_text(encoding='utf-8')) if cached_path.exists() else {'files':[]}
    prior={r['path']:r for r in cached['files']}
    manifest=root/'data/real/source_evaluation.json';sources=json.loads(manifest.read_text(encoding='utf-8')).get('sources',[]) if manifest.exists() else []
    files=[]
    for folder in sorted(p.name for p in (root/'data').iterdir() if p.is_dir() and p.name not in {'intelligence','qc','knowledge'}):
        for p in (root/'data'/folder).rglob('*'):
            if p.is_file():
                item={'path':str(p.resolve()),'relative':str(p.relative_to(root)),'bytes':p.stat().st_size,'role':'derived artifact'}
                if p.suffix=='.parquet':
                    pf=pq.ParquetFile(p);item.update(rows=pf.metadata.num_rows,columns=pf.schema_arrow.names)
                if folder=='raw':item['role']='additional input'
                if p.name=='canonical_inspections.parquet':item['role']='primary real inspection and citation source'
                if p.name=='annual_benchmarks.parquet':item['role']='primary annual templates'
                if folder=='processed' and p.name=='inspections.parquet':item['role']='primary synthetic inspections'
                if p.name=='workspace.db':item['role']='synthetic workspace runs; use distinct original observations, not repeated exports'
                files.append(item)
    candidates={p for p in (root/'data').rglob('*') if p.suffix.lower()=='.xlsx' and not set(p.relative_to(root/'data').parts[:-1]) & {'intelligence','qc','knowledge','observations'}}
    candidates.update(source_paths(root))
    for p in sorted(candidates):
        if not p.exists():files.append({'path':str(p),'role':'missing original','available':False});continue
        try: digest=sha(p)
        except PermissionError:
            files.append({'path':str(p),'role':'source workbook','available':False,'error':'PermissionError reading original; prepared import and prior audit remain available','prior_audit':next((r for r in sources if r['file']==p.name),None)})
            continue
        key=str(p.resolve());old=prior.get(key)
        if old and old.get('sha256')==digest and old.get('sheets'):files.append(old);continue
        item={'path':key,'sha256':digest,'role':'source workbook','available':True,'sheets':[]}
        linked=next((r for r in sources if r['file']==p.name and r['sha256']==digest),None)
        item['covered_by_verified_prepared_import']=bool(linked)
        if linked:
            item['audit']=linked
            item['sheets']=[{'sheet':sheet,'audited_rows':count} for sheet,count in linked.get('sheet_counts',{}).items()]
            item['audited_rows']=linked.get('rows')
            item['note']='SHA-256 matches the audited source for the prepared import; reuse audited counts without rescanning all Excel rows.'
            files.append(item);continue
        wb=load_workbook(p,read_only=True,data_only=True)
        for ws in wb:
            stream=ws.iter_rows(values_only=True);header=next(stream,())
            columns=[str(v) if v is not None else f'column_{i+1}' for i,v in enumerate(header)]
            counts=Counter();ids=set();duplicates=0;n=0
            id_index=next((i for i,k in enumerate(columns) if k.lower()=='inspection id'),None)
            for row in stream:
                if not any(v is not None and str(v).strip() for v in row):continue
                n+=1
                for i,k in enumerate(columns):
                    if i>=len(row) or row[i] is None or str(row[i]).strip()=='':counts[k]+=1
                if id_index is not None and id_index<len(row) and row[id_index] is not None:
                    value=str(row[id_index]);duplicates+=value in ids;ids.add(value)
            item['sheets'].append({'sheet':ws.title,'nonblank_rows_after_first_row':n,'columns':columns,'missing':dict(counts),'distinct_inspection_ids':len(ids) if id_index is not None else None,'repeated_inspection_id_rows':duplicates if id_index is not None else None,'note':'Summary sheets are inventoried as layout rows, not observation counts.'})
        wb.close();files.append(item)
        cached_path.write_text(json.dumps({'files':files},indent=2),encoding='utf-8')
        print('Inventoried',p.name,flush=True)
    report={'files':files,'note':'Prepared mirrors and generated exports are inventoried but not re-ingested as additional observations. Originals are matched by SHA-256 to the audited real import. Summary frequencies remain annual-template data.'}
    cached_path.write_text(json.dumps(report,indent=2),encoding='utf-8');return report
if __name__=='__main__':inventory()
